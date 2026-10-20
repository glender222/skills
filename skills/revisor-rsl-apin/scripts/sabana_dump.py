#!/usr/bin/env python3
"""Vuelca hojas de la sábana de la RSL (.xlsx) a JSON para contrastarlas con el
manuscrito.

Uso:
    python3 sabana_dump.py sabana.xlsx [--sheets "X6 Uso de IA" "X7 PRISMA"] [--json salida.json]

Para cada celda con contenido entrega el valor calculado y, si la celda tiene
fórmula, también la fórmula. Si el valor calculado no está guardado en el
archivo (libro generado por script y nunca abierto en Excel o LibreOffice),
'valor' sale null: esa cifra no es verificable hasta recalcular el libro.
Requiere openpyxl.
"""
from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

try:
    import openpyxl
except ImportError:  # pragma: no cover
    print("Falta openpyxl: pip install openpyxl", file=sys.stderr)
    sys.exit(2)


def dump(path: Path, sheets: list[str] | None) -> dict:
    wb_v = openpyxl.load_workbook(path, data_only=True)
    wb_f = openpyxl.load_workbook(path, data_only=False)
    names = sheets or wb_v.sheetnames
    out = {"archivo": path.name, "hojas": {}, "advertencias": []}
    for name in names:
        if name not in wb_v.sheetnames:
            out["advertencias"].append(f"No existe la hoja '{name}'. Hojas: {wb_v.sheetnames}")
            continue
        wv, wf = wb_v[name], wb_f[name]
        filas = []
        sin_valor = 0
        for row_v, row_f in zip(wv.iter_rows(), wf.iter_rows()):
            celdas = {}
            for cv, cf in zip(row_v, row_f):
                if cf.value is None:
                    continue
                item = {"valor": cv.value if not hasattr(cv.value, "isoformat") else cv.value.isoformat()}
                if isinstance(cf.value, str) and cf.value.startswith("="):
                    item["formula"] = cf.value
                    if cv.value is None:
                        sin_valor += 1
                celdas[cf.coordinate] = item
            if celdas:
                filas.append(celdas)
        out["hojas"][name] = filas
        if sin_valor:
            out["advertencias"].append(
                f"'{name}': {sin_valor} fórmulas sin valor calculado guardado; recalcular el libro antes de contrastar cifras.")
    return out


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("xlsx", type=Path)
    ap.add_argument("--sheets", nargs="*", help="hojas a volcar (por defecto, todas)")
    ap.add_argument("--json", type=Path, help="archivo de salida (por defecto, stdout)")
    a = ap.parse_args(argv)
    data = dump(a.xlsx, a.sheets)
    txt = json.dumps(data, ensure_ascii=False, indent=1, default=str)
    if a.json:
        a.json.write_text(txt, encoding="utf-8")
        print(f"Volcado en {a.json}. Advertencias: {data['advertencias'] or 'ninguna'}")
    else:
        print(txt)
    return 0


if __name__ == "__main__":
    sys.exit(main())
